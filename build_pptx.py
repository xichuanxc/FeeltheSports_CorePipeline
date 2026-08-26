#!/usr/bin/env python3
"""
build_pptx.py — render docs/slides.html into a PowerPoint deck.

    python3 build_pptx.py            # writes docs/FeelingTheTennisGame.pptx

One full-bleed image per slide, rendered from the built HTML deck by headless
Chrome at 2560x1440, plus the five audio clips as clickable media placed exactly
over the buttons they already sit on. Run build_slides.py first; this reads the
built deck, not the source.

WHY IMAGES. The deck is HTML, SVG and web type. Reproducing that as native
PowerPoint shapes would be a second implementation to keep in step with the
first. Rasterising keeps one source of truth; the cost is that text in the pptx
is not selectable or editable, so edits still go through docs/slides.src.html.

WHY THE XML IS PATCHED AFTERWARDS. python-pptx has no audio API: add_movie
writes <a:videoFile> with a video relationship, and no content type is declared
for .wav. PowerPoint may then refuse the file or offer to repair it. The patch
pass rewrites the element, the relationship type, and [Content_Types].xml.

Button geometry is read from the rendered page via getBoundingClientRect rather
than inferred from the CSS, so the audio hit areas line up with what is drawn.
"""

import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

from pptx import Presentation
from pptx.util import Emu
from PIL import Image

DECK = "docs/slides.html"
AUDIO = "docs/audio"
OUT = "docs/FeelingTheTennisGame.pptx"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
RENDER_W, RENDER_H = 2560, 1440
SLIDE_W, SLIDE_H = Emu(12192000), Emu(6858000)      # 13.333 x 7.5in, 16:9


def chrome(args):
    return subprocess.run([CHROME, "--headless", "--disable-gpu",
                           f"--window-size={RENDER_W},{RENDER_H}",
                           "--hide-scrollbars", "--virtual-time-budget=3000"] + args,
                          capture_output=True, text=True)


def advance(n):
    return ('<script>addEventListener("load",function(){for(var k=0;k<%d;k++)'
            'dispatchEvent(new KeyboardEvent("keydown",{key:"ArrowRight"}));' % n)


def render_slides(html, tmp):
    n_slides = html.count('<section class="slide')
    shots = []
    for n in range(n_slides):
        wrap = os.path.join(tmp, f"w{n}.html")
        open(wrap, "w", encoding="utf-8").write(html + advance(n) + "});</script>")
        png = os.path.join(tmp, f"s{n:02d}.png")
        chrome([f"--screenshot={png}", "file://" + wrap])
        if not os.path.exists(png):
            sys.exit(f"chrome failed to render slide {n + 1}")
        shots.append(png)
        print(f"  rendered slide {n + 1}/{n_slides}", end="\r")
    print(f"  rendered {n_slides} slides            ")
    return shots


def clip_rects(html, tmp, one_based):
    """Where the audio buttons actually are, asked of the page itself."""
    js = (advance(one_based - 1) +
          'setTimeout(function(){var o=[];'
          'document.querySelectorAll(".slide.on .clip").forEach(function(b){'
          'var r=b.getBoundingClientRect();'
          'o.push({clip:b.dataset.clip,x:r.x,y:r.y,w:r.width,h:r.height});});'
          'document.body.innerHTML="<pre id=out>"+JSON.stringify(o)+"</pre>";},600);'
          '});</script>')
    wrap = os.path.join(tmp, f"r{one_based}.html")
    open(wrap, "w", encoding="utf-8").write(html + js)
    dom = chrome(["--dump-dom", "file://" + wrap]).stdout
    m = re.search(r'<pre id="out">(\[.*?\])</pre>', dom, re.S)
    return json.loads(m.group(1)) if m else []


def patch_audio(path):
    """python-pptx writes a video shape for a wav; PowerPoint wants an audio one."""
    tmp = path + ".tmp"
    zin, zout = zipfile.ZipFile(path), zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED)
    slides = rels = 0
    for item in zin.infolist():
        data = zin.read(item.filename)
        name = item.filename
        if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
            t = data.decode("utf-8")
            if "a:videoFile" in t:
                t = t.replace("<a:videoFile ", "<a:audioFile ")
                t = t.replace("</a:videoFile>", "</a:audioFile>")
                slides += 1
            data = t.encode("utf-8")
        elif "slides/_rels/" in name:
            t = data.decode("utf-8")
            if "relationships/video" in t:
                t = t.replace("officeDocument/2006/relationships/video",
                              "officeDocument/2006/relationships/audio")
                rels += 1
            data = t.encode("utf-8")
        elif name == "[Content_Types].xml":
            t = data.decode("utf-8")
            if 'Extension="wav"' not in t:
                t = t.replace("</Types>",
                              '<Default Extension="wav" ContentType="audio/x-wav"/></Types>')
            data = t.encode("utf-8")
        zout.writestr(item, data)
    zout.close()
    zin.close()
    shutil.move(tmp, path)
    return slides, rels


def main():
    if not os.path.exists(DECK):
        sys.exit(f"{DECK} not found -- run build_slides.py first")
    html = open(DECK, encoding="utf-8").read()
    secs = re.findall(r'<section class="slide.*?</section>', html, re.S)
    clip_slides = [i + 1 for i, s in enumerate(secs) if 'class="clip"' in s]

    with tempfile.TemporaryDirectory() as tmp:
        shots = render_slides(html, tmp)
        rects = {i: clip_rects(html, tmp, i) for i in clip_slides}

        prs = Presentation()
        prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
        blank = prs.slide_layouts[6]
        ex = lambda px: Emu(int(round(px / RENDER_W * SLIDE_W)))
        ey = lambda px: Emu(int(round(px / RENDER_H * SLIDE_H)))

        for i, png in enumerate(shots, 1):
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(png, 0, 0, SLIDE_W, SLIDE_H)
            for r in rects.get(i, []):
                # poster cut from the slide, so the media shape is invisible and
                # the button the audience sees is the button they click
                poster = os.path.join(tmp, f"poster{i}_{r['clip']}.png")
                Image.open(png).crop((int(r["x"]), int(r["y"]),
                                      int(r["x"] + r["w"]),
                                      int(r["y"] + r["h"]))).save(poster)
                slide.shapes.add_movie(os.path.join(AUDIO, f"{r['clip']}.wav"),
                                       ex(r["x"]), ey(r["y"]), ex(r["w"]), ey(r["h"]),
                                       poster_frame_image=poster, mime_type="audio/wav")
        prs.save(OUT)

    s, rl = patch_audio(OUT)
    print(f"wrote {OUT}  ({len(shots)} slides, "
          f"{os.path.getsize(OUT) / 1024 / 1024:.1f} MB)")
    print(f"  audio on slides {clip_slides}; patched {s} slide parts, {rl} rel parts")
    return 0


if __name__ == "__main__":
    sys.exit(main())
